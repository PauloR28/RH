"""Cadastro de usuários em massa (planilha): modelo, prévia que ignora linhas inválidas e criação das válidas.
Integração contra o banco de DEV (pulada sem banco). Dados de teste: operação `UMASS_xxxx`, e-mails `umassa_*@example.com`."""

from __future__ import annotations

import io
import uuid

import pytest
from openpyxl import Workbook, load_workbook

from _integracao_dev import repositorio_dev
from rh_api.auth import AuthenticatedUser
from rh_api.rbac import ROLE_ADMIN, ROLE_SUPERVISOR, get_role_permissions
from rh_api.repositories.usuarios_massa import ABA_DADOS, COLUNAS

CABECALHO = [c[0] for c in COLUNAS]


def _admin():
    return AuthenticatedUser(username="adm", id_usuario=1, nome="Adm", perfil=ROLE_ADMIN, operacoes=frozenset(), permissions=frozenset(get_role_permissions(ROLE_ADMIN)))


def _planilha(linhas, cabecalho=None):
    wb = Workbook()
    ws = wb.active
    ws.title = ABA_DADOS
    ws.append(cabecalho or CABECALHO)
    for linha in linhas:
        ws.append(linha)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.fixture(scope="module")
def ctx():
    repo = repositorio_dev()
    c = type("Ctx", (), {})()
    c.repo, c.admin = repo, _admin()
    c.op = f"UMASS_{uuid.uuid4().hex[:4].upper()}"
    c.tag = uuid.uuid4().hex[:6]
    repo.upsert_configuration_item("operacoes", {"chave": c.op, "nome": f"Op {c.op}", "ativo": True})
    c.email_sup = f"umassa_sup_{c.tag}@example.com"
    repo.mon_create_usuario(c.admin, {"nome": "Sup", "sobrenome": "Teste", "email": c.email_sup, "perfil": ROLE_SUPERVISOR, "operacoes": [c.op], "provedor_autenticacao": "microsoft"})
    yield c
    conn = repo._connect()
    try:
        cursor = conn.cursor()
        cursor.execute("SELECT id_usuario FROM dbo.usuarios WHERE email LIKE 'umassa_%@example.com'")
        for (id_usuario,) in cursor.fetchall():
            cursor.execute("DELETE FROM dbo.usuarios_supervisores WHERE id_operador = ? OR id_supervisor = ?", (id_usuario, id_usuario))
            cursor.execute("DELETE FROM dbo.usuarios_canais WHERE id_usuario = ?", (id_usuario,))
            cursor.execute("DELETE FROM dbo.usuarios_operacoes WHERE id_usuario = ?", (id_usuario,))
            cursor.execute("DELETE FROM dbo.usuarios_operacoes_historico WHERE id_usuario = ?", (id_usuario,))
            cursor.execute("DELETE FROM dbo.usuarios WHERE id_usuario = ?", (id_usuario,))
        cursor.execute("DELETE FROM dbo.operacoes WHERE chave LIKE 'UMASS_%'")
        conn.commit()
    finally:
        conn.close()


def test_modelo_tem_as_colunas_e_as_listas(ctx):
    wb = load_workbook(io.BytesIO(ctx.repo.usuarios_massa_modelo(ctx.admin)))
    assert wb.sheetnames[0] == ABA_DADOS and [c.value for c in wb[ABA_DADOS][1]] == CABECALHO
    assert wb["Instruções"] and wb["Listas"]["A2"].value


def test_previa_ignora_o_que_esta_incompleto_e_confirmacao_cria_so_os_validos(ctx):
    t = ctx.tag
    nome_op = f"Op {ctx.op}"
    linhas = [
        ["Ana", "Silva", f"umassa_ana_{t}@example.com", "Analista de RH", "RH", "Microsoft", None, None, None, None, None, None],
        ["Beto", "Souza", f"umassa_beto_{t}@example.com", "Operador de Atendimento", "Operador", "Local", "Senha@2026", nome_op, ctx.email_sup, None, None, None],
        ["Caio", "Lima", f"umassa_caio_{t}@example.com", None, "RH", "Microsoft", None, None, None, None, None, None],              # cargo vazio
        ["Duda", "Reis", "email-invalido", "Analista", "RH", "Microsoft", None, None, None, None, None, None],                       # e-mail inválido
        ["Eva", "Melo", f"umassa_eva_{t}@example.com", "Analista", "Perfil Inexistente", "Microsoft", None, None, None, None, None, None],
        ["Fabio", "Nunes", f"umassa_fabio_{t}@example.com", "Analista", "RH", "Local", "curta", None, None, None, None, None],        # senha curta
        ["Gabi", "Dias", f"umassa_ana_{t}@example.com", "Analista", "RH", "Microsoft", None, None, None, None, None, None],          # e-mail repetido
        ["Hugo", "Paz", f"umassa_hugo_{t}@example.com", "Operador", "Operador", "Microsoft", None, nome_op, None, None, None, None],  # operador sem supervisor
        [None] * 12,                                                                                                                  # linha vazia: nem conta
    ]
    arquivo = _planilha(linhas)
    previa = ctx.repo.usuarios_massa_processar(ctx.admin, arquivo)
    assert previa["confirmado"] is False and previa["total_linhas"] == 8
    assert [v["email"] for v in previa["validos"]] == [f"umassa_ana_{t}@example.com", f"umassa_beto_{t}@example.com"]
    assert {i["linha"] for i in previa["ignorados"]} == {4, 5, 6, 7, 8, 9}
    assert "senha" not in str(previa).lower().replace("senha inicial", "")  # a senha nunca volta para a tela
    conn = ctx.repo._connect()
    try:
        c = conn.cursor()
        c.execute("SELECT COUNT(*) FROM dbo.usuarios WHERE email LIKE ?", (f"umassa_%_{t}@example.com",))
        assert c.fetchone()[0] == 1  # só o supervisor do fixture: a prévia não grava nada
    finally:
        conn.close()

    final = ctx.repo.usuarios_massa_processar(ctx.admin, arquivo, confirmar=True)
    assert len(final["criados"]) == 2 and final["falhas"] == []
    beto = next(u for u in ctx.repo.list_system_users(search=f"umassa_beto_{t}") )
    assert "perfil_nome" in beto or "perfil" in beto
    vinculos = ctx.repo.mon_get_vinculos(beto["id_usuario"])
    assert vinculos["operacoes"] == [ctx.op] and len(vinculos["supervisores"]) == 1

    # reenviar o mesmo arquivo: todos já existem e são ignorados
    de_novo = ctx.repo.usuarios_massa_processar(ctx.admin, arquivo)
    assert de_novo["validos"] == [] and any("Já existe" in i["motivo"] for i in de_novo["ignorados"])


def test_arquivo_fora_do_modelo_e_rejeitado(ctx):
    from fastapi import HTTPException

    with pytest.raises(HTTPException) as e:
        ctx.repo.usuarios_massa_processar(ctx.admin, _planilha([["x"]], cabecalho=["Coluna qualquer"]))
    assert e.value.status_code == 422
    with pytest.raises(HTTPException) as e:
        ctx.repo.usuarios_massa_processar(ctx.admin, b"nao e uma planilha")
    assert e.value.status_code == 400
