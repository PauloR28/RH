"""WFM — Histórico (antiga Auditoria): regras puras, paginação por cursor, exportação e fuso. Sem banco."""

from __future__ import annotations

import csv
import io
import sqlite3
from datetime import date, datetime, timedelta, timezone

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pypdf import PdfReader

from rh_api.services import wfm_historico as h

UTC = timezone.utc
BRASIL_MENOS_3 = timezone(timedelta(hours=-3))


# ---------------------------------------------------------------------------
# Nomes das ações e cor da bolinha
# ---------------------------------------------------------------------------
def test_nome_acao_mapeado_e_codigo_bruto_quando_sem_mapeamento():
    assert h.nome_acao("lancar_presenca") == "Lançou presença"
    assert h.nome_acao("solicitar_troca") == "Solicitou troca"
    assert h.nome_acao("acao_que_nao_existe") == "acao_que_nao_existe"
    assert h.nome_acao("") == ""


@pytest.mark.parametrize("acao,depois,tom", [
    ("aprovar_troca", None, "troca"),
    ("solicitar_troca", "{}", "troca"),
    ("lancar_presenca", '{"status": "PRESENTE"}', "presente"),
    ("lancar_presenca", '{"status": "FALTA"}', "falta"),
    ("lancar_presenca", '{"status": "ATESTADO"}', "neutro"),
    ("lancar_presenca", "json quebrado", "neutro"),
    ("publicar_escala", '{"status": "FALTA"}', "neutro"),
])
def test_tom_da_bolinha(acao, depois, tom):
    assert h.tom_acao(acao, depois) == tom


def test_pares_antes_depois_compativel_com_registros_antigos():
    assert h.pares_json(None) == []
    assert h.pares_json('{"status": "PRESENTE", "observacao": ""}') == [
        {"rotulo": "Status", "valor": "PRESENTE"}, {"rotulo": "Observação", "valor": "—"}]
    assert h.pares_json('{"data_ini": "2026-10-09", "ativo": true}') == [
        {"rotulo": "Início", "valor": "09/10/2026"}, {"rotulo": "Ativo", "valor": "Sim"}]
    assert h.pares_json("texto solto de um registro antigo") == [{"rotulo": "", "valor": "texto solto de um registro antigo"}]
    assert h.pares_json('"so uma string"') == [{"rotulo": "", "valor": "so uma string"}]
    assert h.texto_pares(h.pares_json('{"status": "PRESENTE", "observacao": ""}')) == "Status: PRESENTE · Observação: —"


# ---------------------------------------------------------------------------
# Fuso horário
# ---------------------------------------------------------------------------
def test_conversao_de_fuso_servidor_utc_para_brasilia():
    r = h.para_brasilia(datetime(2026, 10, 10, 13, 25, 41), UTC)
    assert (r.hour, r.minute, r.second) == (10, 25, 41)
    assert h.fmt_data(r) == "10/10/2026" and h.fmt_hora(r) == "10:25" and h.fmt_hora(r, segundos=True) == "10:25:41"


def test_conversao_de_fuso_servidor_ja_em_brasilia_nao_desloca():
    r = h.para_brasilia(datetime(2026, 10, 10, 10, 25, 41), BRASIL_MENOS_3)
    assert (r.hour, r.minute) == (10, 25)


def test_conversao_de_fuso_virada_de_dia():
    r = h.para_brasilia(datetime(2026, 10, 11, 1, 30), UTC)  # 01:30 UTC = 22:30 do dia anterior em Brasília
    assert h.fmt_data(r) == "10/10/2026" and h.fmt_hora(r) == "22:30"


def test_limites_do_periodo_em_hora_local_do_servidor():
    ini, fim = h.limites_periodo(date(2026, 10, 1), date(2026, 10, 31), UTC)
    assert ini == datetime(2026, 10, 1, 3, 0) and fim == datetime(2026, 11, 1, 3, 0)  # meia-noite de Brasília = 03:00 UTC
    ini, fim = h.limites_periodo(date(2026, 10, 1), None, BRASIL_MENOS_3)
    assert ini == datetime(2026, 10, 1, 0, 0) and fim is None


# ---------------------------------------------------------------------------
# Filtros combinados
# ---------------------------------------------------------------------------
def test_filtros_combinados_geram_condicoes_e_parametros_na_mesma_ordem():
    f = h.FiltroHistorico(texto="ana", operacao="TI", entidade="presenca", operador="Ana Souza", supervisor="Marcos Lima",
                          data_ini=date(2026, 10, 1), data_fim=date(2026, 10, 31))
    where, params = h.montar_where(f, permitidas=None, fuso_servidor=UTC)
    assert where.count("?") == len(params)
    for trecho in ("a.operacao = ?", "a.entidade = ?", "a.usuario_nome = ?", "s.nome = ?", "a.criado_em >= ?", "a.criado_em < ?", "LIKE ?"):
        assert trecho in where
    assert params[:4] == ["TI", "presenca", "Ana Souza", "Marcos Lima"]
    assert params[4] == datetime(2026, 10, 1, 3) and params[5] == datetime(2026, 11, 1, 3)
    assert params[6:] == ["%ana%"] * 5


def test_escopo_de_operacao_inclui_chaves_virtuais_e_sem_acesso_nao_retorna_nada():
    where, params = h.montar_where(h.FiltroHistorico(), permitidas=["TI", "RH"])
    assert "a.operacao IN (?,?)" in where and params == ["TI", "RH", "TI::%", "RH::%"]
    assert h.montar_where(h.FiltroHistorico(), permitidas=[]) == ("1 = 0", [])
    assert h.montar_where(h.FiltroHistorico(), permitidas=None) == ("1 = 1", [])


def test_texto_livre_escapa_curingas():
    _, params = h.montar_where(h.FiltroHistorico(texto="50%_[x]"), permitidas=None)
    assert params[0] == "%50\\%\\_\\[x]%"


def test_escopo_selecionados_usa_so_ids_e_escopo_de_operacao():
    f = h.FiltroHistorico(operacao="TI", texto="x", operador="y")
    where, params = h.montar_where(f, permitidas=["TI"], ids=[7, 9])
    assert "a.id_auditoria IN (?,?)" in where and "usuario_nome" not in where and "a.operacao = ?" not in where
    assert params[-2:] == [7, 9]


# ---------------------------------------------------------------------------
# Paginação por cursor: o predicado real, executado num banco de teste, não repete nem pula linha
# ---------------------------------------------------------------------------
def _banco_com(n: int) -> sqlite3.Connection:
    con = sqlite3.connect(":memory:")
    con.execute("CREATE TABLE wfm_auditoria (id_auditoria INTEGER PRIMARY KEY AUTOINCREMENT, criado_em TEXT)")
    base = datetime(2026, 10, 10, 10, 0, 0)
    # vários registros no MESMO instante (desempate pelo id) e outros em instantes distintos
    for i in range(n):
        con.execute("INSERT INTO wfm_auditoria (criado_em) VALUES (?)", ((base + timedelta(seconds=i // 3)).isoformat(),))
    return con


def _pagina(con, cursor_pos, tamanho):
    sql = "SELECT a.id_auditoria, a.criado_em FROM wfm_auditoria a WHERE 1 = 1"
    params: list = []
    if cursor_pos:
        sql += f" AND {h.PREDICADO_CURSOR}"
        c, i = cursor_pos
        params += [c.isoformat(), c.isoformat(), i]
    sql += f" ORDER BY {h.ORDEM} LIMIT {tamanho + 1}"
    return con.execute(sql, params).fetchall()


def test_cursor_percorre_tudo_sem_repetir_nem_pular_mesmo_com_registros_novos_no_meio():
    con = _banco_com(50)
    vistos, token, tamanhos = [], None, [8, 20, 20, 20]
    for passo, tamanho in enumerate(tamanhos):
        linhas = _pagina(con, h.decodificar_cursor(token), tamanho)
        pagina, tem_mais = linhas[:tamanho], len(linhas) > tamanho
        vistos += [r[0] for r in pagina]
        if passo == 0:  # entram registros novos (mais recentes) entre a 1ª e a 2ª página
            for _ in range(5):
                con.execute("INSERT INTO wfm_auditoria (criado_em) VALUES (?)", (datetime(2026, 10, 11).isoformat(),))
        token = h.codificar_cursor(datetime.fromisoformat(pagina[-1][1]), pagina[-1][0]) if tem_mais else None
        if not tem_mais:
            break
    assert len(vistos) == len(set(vistos)) == 50          # sem repetir
    assert sorted(vistos) == list(range(1, 51))             # sem pular (os 5 novos ficam para o "recarregar")
    assert vistos == sorted(vistos, key=lambda i: (-(i - 1) // 3, -i))  # mais recentes primeiro, desempate por id


def test_cursor_invalido_responde_400():
    with pytest.raises(HTTPException) as erro:
        h.decodificar_cursor("isto-nao-e-cursor")
    assert erro.value.status_code == 400
    assert h.decodificar_cursor("") is None and h.decodificar_cursor(None) is None
    quando = datetime(2026, 10, 10, 10, 25, 41, 123000)
    assert h.decodificar_cursor(h.codificar_cursor(quando, 42)) == (quando, 42)


# ---------------------------------------------------------------------------
# Item, painel da troca e exportação
# ---------------------------------------------------------------------------
def _linha(**extra):
    base = {"id_auditoria": 10, "operacao": "TI", "usuario_nome": "Ana Souza", "perfil": "operador", "acao": "lancar_presenca",
            "entidade": "presenca", "entidade_id": "42:2026-10-09", "antes_json": None, "depois_json": '{"status": "PRESENTE", "observacao": ""}',
            "justificativa": None, "criado_em": datetime(2026, 10, 10, 13, 25, 41), "dia_plantao": date(2026, 10, 9),
            "hora_inicio": "08:00", "hora_fim": "17:00", "supervisor_nome": "Marcos Lima"}
    base.update(extra)
    return base


def test_item_de_presenca_traz_plantao_supervisor_e_registro_em_brasilia():
    item = h.montar_item(_linha(), fuso_servidor=UTC)
    assert (item["data"], item["hora_inicio"], item["hora_fim"], item["supervisor"]) == ("09/10/2026", "08:00", "17:00", "Marcos Lima")
    assert item["registrado_em"] == "10/10/2026 · 10:25:41"
    assert (item["acao_nome"], item["acao_tom"], item["eh_troca"], item["troca"]) == ("Lançou presença", "presente", False, None)
    assert item["depois_texto"] == "Status: PRESENTE · Observação: —" and item["antes"] == []


def test_registro_antigo_sem_plantao_mostra_vazio_e_usa_a_data_do_registro():
    item = h.montar_item(_linha(entidade="escala", entidade_id="2026-10", acao="publicar_escala", dia_plantao=None, hora_inicio=None, hora_fim=None, supervisor_nome=None), fuso_servidor=UTC)
    assert item["data"] == "10/10/2026" and item["hora_inicio"] == "" and item["supervisor"] == ""
    assert h.linha_exportacao(item)[1:3] == ["—", "—"]


def test_painel_da_troca_le_solicitou_aceitou_e_supervisor_dos_eventos():
    eventos = [
        {"evento": "SOLICITADA", "por_nome": "Ana Souza", "criado_em": datetime(2026, 10, 10, 13, 12)},
        {"evento": "AGUARDANDO_APROVACAO", "por_nome": "Bruno Reis", "criado_em": datetime(2026, 10, 10, 14, 0)},
        {"evento": "APROVADA", "por_nome": "Marcos Lima", "criado_em": datetime(2026, 10, 10, 15, 30)},
    ]
    item = h.montar_item(_linha(acao="aprovar_troca", entidade="troca", entidade_id="5"), eventos, UTC)
    assert item["eh_troca"] and item["acao_tom"] == "troca"
    assert item["troca"]["solicitou"] == {"nome": "Ana Souza", "data": "10/10/2026", "hora": "10:12", "em": item["troca"]["solicitou"]["em"]}
    assert (item["troca"]["aceitou"]["nome"], item["troca"]["aceitou"]["hora"]) == ("Bruno Reis", "11:00")
    assert (item["troca"]["efetivou"]["nome"], item["troca"]["efetivou"]["hora"]) == ("Marcos Lima", "12:30")
    # troca ainda sem aceite: só o primeiro passo existe
    parcial = h.montar_item(_linha(acao="solicitar_troca", entidade="troca", entidade_id="6"), eventos[:1], UTC)
    assert parcial["troca"]["aceitou"] is None and parcial["troca"]["efetivou"] is None


def _linhas_exemplo(n=3):
    eventos = [{"evento": "SOLICITADA", "por_nome": "Ana", "criado_em": datetime(2026, 10, 10, 13, 12)}]
    return [h.linha_exportacao(h.montar_item(_linha(id_auditoria=i, entidade="troca", entidade_id=str(i), acao="solicitar_troca") if i == 1
                                             else _linha(id_auditoria=i), eventos if i == 1 else None, UTC)) for i in range(1, n + 1)]


def test_exportacao_xlsx_cabecalhos_em_portugues_e_contagem_de_linhas():
    conteudo = h.gerar_xlsx_fluxo(_linhas_exemplo(3))
    folha = load_workbook(io.BytesIO(conteudo)).active
    cabecalhos = [c.value for c in folha[1]]
    assert cabecalhos == [rotulo for _, rotulo in h.COLUNAS_EXPORTACAO]
    assert {"Data", "Hora início", "Hora final", "Quem solicitou", "Quem aceitou", "Supervisor que trocou", "Justificativa", "Antes", "Depois", "Registrado em (data)", "Registrado em (hora)"} <= set(cabecalhos)
    assert folha.max_row == 1 + 3
    assert folha.cell(row=2, column=cabecalhos.index("Quem solicitou") + 1).value == "Ana"


def test_exportacao_csv_em_fluxo_tem_bom_cabecalhos_e_linhas():
    pedacos = list(h.gerar_csv_fluxo(iter(_linhas_exemplo(1200)), tamanho_bloco=500))
    assert len(pedacos) >= 3  # em blocos, não de uma vez
    texto = b"".join(pedacos).decode("utf-8")
    assert texto.startswith("﻿")
    linhas = list(csv.reader(io.StringIO(texto.lstrip("﻿")), delimiter=";"))
    assert linhas[0] == [rotulo for _, rotulo in h.COLUNAS_EXPORTACAO] and len(linhas) == 1 + 1200


def test_exportacao_protege_celulas_que_parecem_formula():
    linha = _linhas_exemplo(1)[0]
    linha[h.COLUNAS_EXPORTACAO.index(("justificativa", "Justificativa"))] = "=HYPERLINK(\"x\")"
    folha = load_workbook(io.BytesIO(h.gerar_xlsx_fluxo([linha]))).active
    assert str(folha.cell(row=2, column=len(h.COLUNAS_EXPORTACAO)).value).startswith("'")


def test_exportacao_pdf_e_valido_e_repete_cabecalho_por_pagina():
    pdf = h.gerar_pdf(iter(_linhas_exemplo(1) * 400))
    leitor = PdfReader(io.BytesIO(pdf))
    assert len(leitor.pages) > 1
    assert "Quem solicitou" in leitor.pages[1].extract_text() or "solicitou" in leitor.pages[1].extract_text()


def test_escopo_selecionados_respeita_os_identificadores_e_o_limite():
    assert h.validar_exportacao("XLSX", "selecionados", [3, 1, 3]) == ("xlsx", "selecionados", [1, 3])
    assert h.validar_exportacao("csv", "filtro", [1, 2]) == ("csv", "filtro", None)  # ids ignorados no escopo "filtro"
    for formato, escopo, ids in (("doc", "filtro", None), ("xlsx", "outro", None), ("xlsx", "selecionados", []), ("xlsx", "selecionados", list(range(h.LIMITE_IDS + 1)))):
        with pytest.raises(HTTPException):
            h.validar_exportacao(formato, escopo, ids)
    h.validar_limite_linhas(h.LIMITE_LINHAS, "xlsx")
    with pytest.raises(HTTPException) as erro:
        h.validar_limite_linhas(h.LIMITE_LINHAS + 1, "csv")
    assert erro.value.status_code == 413
    with pytest.raises(HTTPException):
        h.validar_limite_linhas(h.LIMITE_LINHAS_PDF + 1, "pdf")


def test_nome_do_arquivo_usa_data_e_hora_de_brasilia():
    assert h.nome_arquivo("xlsx", datetime(2026, 10, 10, 13, 25, tzinfo=UTC)) == "historico_2026-10-10_1025.xlsx"
