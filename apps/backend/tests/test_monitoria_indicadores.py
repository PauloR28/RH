"""Indicadores, exportação e tema da Monitoria (testes puros)."""

from __future__ import annotations

import io
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

API_DIR = Path(__file__).resolve().parents[1]
if str(API_DIR) not in sys.path:
    sys.path.insert(0, str(API_DIR))

from rh_api.services import monitoria_indicadores as ind
from rh_api.services.monitoria_export import gerar_csv, gerar_xlsx, proteger_celula
from rh_api.services.monitoria_tema import contraste, derivar_tokens, normalizar_hex, validar_cor_primaria


def _l(id_operador, nome, nota, data="2026-09-10", *, anulada=False, resultado="", ncg=False, status="FINALIZADA", respostas=None, pilares=None, **extra):
    return {
        "id_operador": id_operador, "operador_nome": nome, "operacao": "CRF", "operacao_nome": "CRF", "equipe_nome": "A",
        "avaliador_nome": "Q", "data": date.fromisoformat(data), "nota": nota, "anulada": anulada, "resultado": resultado,
        "possui_ncg": ncg, "status": status, "respostas": respostas or [], "pilares": pilares or {"conhecimento": {}, "encantamento": {}},
        **extra,
    }


def _resp(bloco, criterio, resposta, peso=5.0, status="AVALIADO"):
    return {"id_bloco": bloco, "bloco_nome": bloco.upper(), "bloco_status": status, "id_criterio": criterio, "pergunta": criterio.upper(),
            "peso": peso, "resposta": resposta}


# ---------------------------------------------------------------------------
# Média, válidas x realizadas, NCG, anuladas
# ---------------------------------------------------------------------------
def test_nota_media_usa_so_validas_e_ncg_entra_com_zero():
    linhas = [_l(1, "Ana", 90), _l(1, "Ana", 80), _l(1, "Ana", 0, ncg=True), _l(1, "Ana", 100, anulada=True),
              _l(1, "Ana", 55, resultado="ANULADA")]
    r = ind.resumo(linhas)
    assert r["quantidade_realizadas"] == 5 and r["quantidade_validas"] == 3 and r["quantidade_anuladas"] == 2
    assert r["nota_media"] == 56.67  # (90 + 80 + 0) / 3 — anuladas fora, NCG dentro com 0
    assert r["ncg"] == 1 and r["melhor_nota"] == 90


def test_sem_monitorias_validas_nao_quebra():
    r = ind.resumo([_l(1, "Ana", 70, anulada=True)])
    assert r["nota_media"] is None and r["faixa"] is None and r["quantidade_validas"] == 0


def test_evolucao_so_validas_por_dia_semana_e_mes():
    linhas = [_l(1, "A", 80, "2026-09-07"), _l(1, "A", 100, "2026-09-08"), _l(1, "A", 10, "2026-09-08", anulada=True), _l(1, "A", 70, "2026-10-01")]
    assert ind.evolucao(linhas, "dia") == [
        {"periodo": "2026-09-07", "nota_media": 80.0, "quantidade": 1}, {"periodo": "2026-09-08", "nota_media": 100.0, "quantidade": 1},
        {"periodo": "2026-10-01", "nota_media": 70.0, "quantidade": 1}]
    assert ind.evolucao(linhas, "semana")[0] == {"periodo": "2026-09-07", "nota_media": 90.0, "quantidade": 2}
    assert [e["periodo"] for e in ind.evolucao(linhas, "mes")] == ["2026-09", "2026-10"]


# ---------------------------------------------------------------------------
# Top N: desempate por quantidade válida e depois por nome
# ---------------------------------------------------------------------------
def test_top_n_com_desempate_e_opcoes_permitidas():
    linhas = (
        [_l(1, "Bruno", 90), _l(1, "Bruno", 90)]
        + [_l(2, "Carla", 90)]
        + [_l(3, "Amanda", 90), _l(3, "Amanda", 90)]
        + [_l(4, "Diego", 100, anulada=True)]      # anulada: fora do ranking
        + [_l(5, "Eva", 95)]
    )
    top = ind.top_n(linhas, 3)
    assert [t["operador_nome"] for t in top] == ["Eva", "Amanda", "Bruno"]  # 95; empate 90: mais válidas (2), depois nome
    assert top[0]["posicao"] == 1 and top[1]["quantidade_validas"] == 2
    assert all(t["operador_nome"] != "Diego" for t in ind.top_n(linhas, 15))
    assert len(ind.top_n(linhas, 7)) == 4  # 7 não é opção válida → usa 5 (só 4 elegíveis)
    assert ind.TOPS_PERMITIDOS == (3, 5, 10, 15)


def test_operador_com_uma_unica_monitoria_pode_aparecer_no_top():
    assert ind.top_n([_l(1, "Solo", 99)], 5)[0]["quantidade_validas"] == 1


# ---------------------------------------------------------------------------
# Critérios com maior erro: N/A fora do denominador; NÃO ≠ NCG
# ---------------------------------------------------------------------------
def test_criterios_com_maior_erro_ignora_na_e_separa_ncg():
    linhas = [
        _l(1, "A", 80, respostas=[_resp("b1", "c1", "NAO"), _resp("b1", "c2", "SIM")]),
        _l(2, "B", 90, respostas=[_resp("b1", "c1", "SIM"), _resp("b1", "c2", "NA")]),
        _l(3, "C", 0, ncg=True, respostas=[_resp("b1", "c1", "NCG"), _resp("b1", "c2", "NA")]),
        _l(4, "D", 100, anulada=True, respostas=[_resp("b1", "c1", "NAO")]),
    ]
    erros = {e["id_criterio"]: e for e in ind.criterios_com_maior_erro(linhas)}
    assert erros["c1"]["avaliados"] == 3 and erros["c1"]["nao"] == 1 and erros["c1"]["ncg"] == 1
    assert erros["c1"]["pct_erro"] == 33.33 and erros["c1"]["pct_ncg"] == 33.33
    assert "c2" not in erros  # 1 SIM + 2 N/A: nenhum erro


def test_desempenho_por_bloco_ignora_blocos_nulos():
    linhas = [
        _l(1, "A", 80, respostas=[_resp("b1", "c1", "SIM"), _resp("b1", "c2", "NAO")]),
        _l(2, "B", 80, respostas=[_resp("b1", "c1", "SIM", status="NULO"), _resp("b1", "c2", "NA", status="NULO")]),
    ]
    assert ind.desempenho_por_bloco(linhas) == [{"id_bloco": "b1", "bloco": "B1", "percentual_medio": 50.0, "quantidade": 1}]


def test_pilares_independentes_da_nota():
    linhas = [_l(1, "A", 10, pilares={"conhecimento": {"Agilidade": 8, "Segurança": 6}, "encantamento": {"Empatia": 10}}),
              _l(2, "B", 100, pilares={"conhecimento": {"Agilidade": 10, "Segurança": 8}, "encantamento": {"Empatia": 6}}),
              _l(3, "C", 90, anulada=True, pilares={"conhecimento": {"Agilidade": 1}, "encantamento": {}})]
    p = ind.pilares(linhas)
    assert p["conhecimento"]["indicadores"] == {"Agilidade": 9.0, "Segurança": 7.0} and p["conhecimento"]["media"] == 8.0
    assert p["encantamento"]["media"] == 8.0


def test_pendencias_e_feedback_vencido():
    agora = datetime(2026, 9, 20, 12, 0)
    linhas = [
        _l(1, "A", 80, status="FEEDBACK_PENDENTE", sla_limite=agora - timedelta(hours=1)),
        _l(2, "B", 80, status="FEEDBACK_PENDENTE", sla_limite=agora + timedelta(hours=5)),
        _l(3, "C", 80, status="AGUARDANDO_CONFIRMACAO", feedback_aplicado=True),
        _l(4, "D", 80, status="REANALISE", feedback_aplicado=True),
        _l(5, "E", 80, status="FINALIZADA", feedback_aplicado=True),
    ]
    p = ind.pendencias(linhas, agora)
    assert p == {"feedbacks_pendentes": 2, "feedbacks_aplicados": 3, "confirmacoes_e_contestacoes_pendentes": 1,
                 "baixas_ou_confirmacoes_pendentes": 1, "feedbacks_vencidos": 1}


def test_agrupar_por_dimensao_ordena_por_nota():
    linhas = [_l(1, "Ana", 80), _l(2, "Bia", 95), _l(2, "Bia", 97)]
    g = ind.agrupar(linhas, "operador")
    assert [x["rotulo"] for x in g] == ["Bia", "Ana"] and g[0]["nota_media"] == 96.0 and g[0]["quantidade_validas"] == 2


def test_distribuicao_usa_o_rotulo_guardado_na_monitoria():
    linhas = [_l(1, "A", 92, nivel="Faixa Antiga"), _l(2, "B", 92)]
    labels = {d["label"] for d in ind.distribuicao_por_faixa(linhas)}
    assert labels == {"Faixa Antiga", "Muito bom"}


# ---------------------------------------------------------------------------
# Exportação segura (XLSX/CSV) — mesma base de linhas
# ---------------------------------------------------------------------------
def test_protecao_contra_injecao_de_formula():
    for perigoso in ("=1+1", "+cmd", "-2", "@SUM(A1)", "\tTAB"):
        assert proteger_celula(perigoso).startswith("'")
    assert proteger_celula("Normal") == "Normal" and proteger_celula(10) == 10


def test_csv_e_xlsx_gravam_os_mesmos_dados_com_protecao():
    colunas, linhas = ["ID", "OPERADOR", "NOTA", "NCG"], [["12345678", "=HYPERLINK(\"x\")", 91.5, True], ["87654321", "Ana", None, False]]
    texto = gerar_csv(colunas, linhas).decode("utf-8-sig")
    assert "'=HYPERLINK" in texto and "Sim" in texto and texto.splitlines()[0] == "ID;OPERADOR;NOTA;NCG"
    wb = load_workbook(io.BytesIO(gerar_xlsx([{"nome": "Monitorias", "colunas": colunas, "linhas": linhas}])))
    ws = wb["Monitorias"]
    assert [c.value for c in ws[1]] == colunas
    assert ws["B2"].value.startswith("'=") and ws["C2"].value == 91.5 and ws["D2"].value == "Sim" and ws["C3"].value in (None, "")


# ---------------------------------------------------------------------------
# Tema por operação: contraste e tokens
# ---------------------------------------------------------------------------
def test_tema_recusa_cor_sem_contraste_e_deriva_tokens():
    assert normalizar_hex("#ABC") == "#aabbcc"
    cor, erros = validar_cor_primaria("#0a4b8c")
    assert cor == "#0a4b8c" and erros == []
    _, erros = validar_cor_primaria("#ffe600")  # amarelo: botão com texto branco ilegível
    assert erros and "contraste" in erros[0]
    assert validar_cor_primaria("azul")[1]
    tokens = derivar_tokens("#0a4b8c")
    assert set(tokens) == {"brand", "brand-ink", "brand-soft", "brand-dark-mode"} and tokens["brand"] == "#0a4b8c"
    assert contraste("#000000", "#ffffff") == 21.0
