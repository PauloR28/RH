"""Geração de XLSX/CSV da Monitoria (promt.txt §5.19–5.20).

Sempre parte das MESMAS linhas produzidas pelo motor de indicadores/relatórios
(nunca recalcula nada). Protege contra injeção de fórmulas: qualquer célula de
texto que comece com = + - @ (ou tab/CR) recebe apóstrofo antes de ser gravada."""

from __future__ import annotations

import csv
import io
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_PERIGOSOS = ("=", "+", "-", "@", "\t", "\r")


def proteger_celula(valor: Any) -> Any:
    if isinstance(valor, str) and valor.startswith(_PERIGOSOS):
        return "'" + valor
    return valor


def _valor(valor: Any) -> Any:
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "Sim" if valor else "Não"
    return proteger_celula(valor)


def gerar_csv(colunas: list[str], linhas: Iterable[list[Any]]) -> bytes:
    saida = io.StringIO()
    escritor = csv.writer(saida, delimiter=";", lineterminator="\r\n")
    escritor.writerow([proteger_celula(c) for c in colunas])
    for linha in linhas:
        escritor.writerow([_valor(v) for v in linha])
    return ("﻿" + saida.getvalue()).encode("utf-8")  # BOM: abre acentuado no Excel


def gerar_xlsx(abas: list[dict[str, Any]]) -> bytes:
    """`abas` = [{"nome": str, "colunas": [...], "linhas": [[...]]}]."""
    wb = Workbook()
    wb.remove(wb.active)
    for aba in abas:
        ws = wb.create_sheet(str(aba["nome"])[:31] or "Dados")
        ws.append([proteger_celula(c) for c in aba["colunas"]])
        for celula in ws[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="0A4B8C")
            celula.alignment = Alignment(vertical="center")
        for linha in aba["linhas"]:
            ws.append([_valor(v) for v in linha])
        for indice, coluna in enumerate(aba["colunas"], start=1):
            largura = max([len(str(coluna))] + [len(str(l[indice - 1])) for l in aba["linhas"][:200] if indice - 1 < len(l)])
            ws.column_dimensions[get_column_letter(indice)].width = min(60, max(10, largura + 2))
        ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def sim_nao(valor: Any) -> str:
    return "Sim" if valor else "Não"
