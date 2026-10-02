"""Cadastro de usuários em massa (Configurações > Usuários > Ações).

Fluxo: o usuário baixa a planilha modelo, preenche uma linha por usuário e envia o arquivo. O servidor LÊ e ANALISA a
planilha (prévia): linhas completas e válidas viram "usuários novos a incluir"; linhas incompletas, inválidas ou que não
puderam ser interpretadas são IGNORADAS com o motivo. Ao confirmar, o mesmo arquivo é reenviado e reanalisado (nada fica
guardado entre as etapas, e a senha nunca volta para a tela) e só as linhas válidas são criadas, uma a uma: a falha de uma
linha não desfaz as demais.

Regras por linha = as do formulário de usuário (mesmas validações de e-mail duplicado, vínculos da Monitoria, escopo de
operação e nível de perfil que o ator pode gerir). Mixin do DatabaseRepository.
"""

from __future__ import annotations

import io
import re
import unicodedata
from typing import Any

from fastapi import HTTPException, status
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..rbac import ROLE_ADMIN, ROLE_CONTROL_DESK, ROLE_DEFINITIONS, ROLE_OPERATOR, ROLE_QUALIDADE, ROLE_SUPERVISOR
from ..services.helpers import normalize_text
from ..services.monitoria_scope import pode_gerenciar_perfil, pode_ver_operacao, validar_vinculos

MAX_LINHAS = 500
MAX_BYTES = 2 * 1024 * 1024
ABA_DADOS = "Usuários"
ABA_AJUDA = "Instruções"
ABA_LISTAS = "Listas"
_EMAIL = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
_PERFIS_MONITORIA = (ROLE_OPERATOR, ROLE_SUPERVISOR, ROLE_QUALIDADE, ROLE_CONTROL_DESK)

# (rótulo, chave interna, obrigatória sempre?, dica)
COLUNAS = (
    ("Nome", "nome", True, "Primeiro nome."),
    ("Sobrenome", "sobrenome", True, "Sobrenome."),
    ("E-mail", "email", True, "E-mail corporativo; também será o login."),
    ("Cargo", "cargo", True, "Ex.: Analista de RH Pleno."),
    ("Perfil", "perfil", True, "Escolha um perfil da lista."),
    ("Tipo de acesso", "acesso", True, "Microsoft ou Local."),
    ("Senha inicial", "senha", False, "Obrigatória só no acesso Local (mín. 8 caracteres). Deixe em branco no acesso Microsoft."),
    ("Operação", "operacao", False, "Nome ou chave da operação. Obrigatória para Operador, Supervisor e Qualidade; Supervisor/Qualidade aceitam várias, separadas por ;."),
    ("Supervisor responsável", "supervisor", False, "E-mail do supervisor (até 2, separados por ;). Obrigatório só para Operador."),
    ("Equipe", "equipe", False, "Opcional (Operador). Nome de uma equipe da operação."),
    ("Turno", "turno", False, "Opcional (Operador e Supervisor). Turno do catálogo."),
    ("Canais", "canais", False, "Opcional (Operador e Supervisor). Canais da operação, separados por ;."),
)


def _norm(valor: Any) -> str:
    texto = unicodedata.normalize("NFD", str(valor if valor is not None else ""))
    return re.sub(r"\s+", " ", "".join(c for c in texto if unicodedata.category(c) != "Mn")).strip().lower()


def _texto(valor: Any) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return re.sub(r"\s+", " ", str(valor)).strip()


def _lista(valor: str) -> list[str]:
    return [parte.strip() for parte in re.split(r"[;\n]", valor or "") if parte.strip()]


def _perfis_visiveis() -> list:
    return [r for r in ROLE_DEFINITIONS.values() if not r.hidden]


class UsuariosMassaRepositoryMixin:
    # ------------------------------------------------------------------
    # Planilha modelo
    # ------------------------------------------------------------------
    def usuarios_massa_modelo(self, actor) -> bytes:
        conn = self._connect()
        try:
            cursor = conn.cursor()
            cursor.execute("SELECT chave, nome FROM dbo.operacoes WHERE ativo = 1 ORDER BY nome")
            operacoes = [(normalize_text(r[0]), normalize_text(r[1])) for r in cursor.fetchall() if pode_ver_operacao(actor.perfil, actor.operacoes, normalize_text(r[0]))]
        finally:
            conn.close()
        turnos = [t["valor"] for t in self.mon_list_catalogo("turno")]
        perfis = [r.name for r in _perfis_visiveis() if r.id != ROLE_ADMIN or actor.perfil == ROLE_ADMIN]

        wb = Workbook()
        ws = wb.active
        ws.title = ABA_DADOS
        ws.append([c[0] for c in COLUNAS])
        for celula in ws[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="0A4B8C")
            celula.alignment = Alignment(vertical="center", wrap_text=True)
        for i, (rotulo, _, _, _) in enumerate(COLUNAS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(16, len(rotulo) + 6)
        ws.freeze_panes = "A2"
        for i in range(2, MAX_LINHAS + 2):
            ws.cell(row=i, column=7).number_format = "@"  # senha como texto
            ws.cell(row=i, column=9).number_format = "@"

        listas = wb.create_sheet(ABA_LISTAS)
        listas.append(["Perfis", "Tipos de acesso", "Operações (nome)", "Operações (chave)", "Turnos"])
        for i in range(max(len(perfis), 2, len(operacoes), len(turnos))):
            listas.append([
                perfis[i] if i < len(perfis) else None,
                ["Microsoft", "Local"][i] if i < 2 else None,
                operacoes[i][1] if i < len(operacoes) else None,
                operacoes[i][0] if i < len(operacoes) else None,
                turnos[i] if i < len(turnos) else None,
            ])
        for celula in listas[1]:
            celula.font = Font(bold=True)
        for i in range(1, 6):
            listas.column_dimensions[get_column_letter(i)].width = 26
        dv_perfil = DataValidation(type="list", formula1=f"={ABA_LISTAS}!$A$2:$A${len(perfis) + 1}", allow_blank=True)
        dv_acesso = DataValidation(type="list", formula1=f"={ABA_LISTAS}!$B$2:$B$3", allow_blank=True)
        for dv, col in ((dv_perfil, "E"), (dv_acesso, "F")):
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{MAX_LINHAS + 1}")

        ajuda = wb.create_sheet(ABA_AJUDA)
        ajuda.append(["Como preencher"])
        ajuda["A1"].font = Font(bold=True, size=14)
        for linha in (
            f"Preencha a aba \"{ABA_DADOS}\": uma linha por usuário, até {MAX_LINHAS} linhas. Não altere os títulos das colunas.",
            "Todas as colunas devem ser preenchidas, exceto as marcadas como condicionais/opcionais abaixo.",
            "Linhas incompletas ou que o Conecta não conseguir interpretar são ignoradas (você vê o motivo antes de confirmar).",
            "O login de cada usuário será o próprio e-mail. No acesso Local, a pessoa troca a senha inicial no primeiro acesso.",
            "Consulte a aba \"Listas\" para os valores aceitos (perfis, operações e turnos).",
            "",
        ):
            ajuda.append([linha])
        ajuda.append(["Coluna", "Preenchimento"])
        for celula in ajuda[ajuda.max_row]:
            celula.font = Font(bold=True)
        for rotulo, _, obrigatoria, dica in COLUNAS:
            ajuda.append([rotulo, ("Obrigatória. " if obrigatoria else "") + dica])
        ajuda.append([])
        ajuda.append(["Exemplo", "Maria | Souza | maria.souza@empresa.com.br | Operadora de Atendimento | Operador | Local | Senha@2026 | (operação) | supervisor@empresa.com.br"])
        ajuda.column_dimensions["A"].width = 28
        ajuda.column_dimensions["B"].width = 110
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    # ------------------------------------------------------------------
    # Leitura e análise
    # ------------------------------------------------------------------
    @staticmethod
    def _um_ler_planilha(conteudo: bytes) -> list[tuple[int, dict[str, str]]]:
        if not conteudo or len(conteudo) > MAX_BYTES:
            raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE if conteudo else status.HTTP_400_BAD_REQUEST,
                                detail="Envie uma planilha .xlsx de até 2 MB.")
        if conteudo[:2] != b"PK":
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Arquivo inválido: envie a planilha .xlsx baixada do Conecta.")
        try:
            wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
            ws = wb[ABA_DADOS] if ABA_DADOS in wb.sheetnames else wb.worksheets[0]
            linhas = list(ws.iter_rows(values_only=True, max_row=MAX_LINHAS + 2))
        except Exception as exc:  # noqa: BLE001 - qualquer falha de leitura = arquivo que não é uma planilha válida
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Não foi possível ler a planilha. Use o modelo baixado do Conecta.") from exc
        if not linhas:
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="A planilha está vazia.")
        cabecalho = [_norm(c) for c in linhas[0]]
        indices = {}
        for rotulo, chave, _, _ in COLUNAS:
            if _norm(rotulo) in cabecalho:
                indices[chave] = cabecalho.index(_norm(rotulo))
        faltando = [r for r, k, _, _ in COLUNAS if k not in indices]
        if faltando:
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                                detail=f"Planilha fora do modelo: faltam as colunas {', '.join(faltando)}. Baixe o modelo novamente.")
        saida = []
        for numero, linha in enumerate(linhas[1:], start=2):
            dados = {k: _texto(linha[i]) if i < len(linha) else "" for k, i in indices.items()}
            if any(dados.values()):
                saida.append((numero, dados))
        if len(saida) > MAX_LINHAS:
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=f"A planilha tem mais de {MAX_LINHAS} usuários. Divida em arquivos menores.")
        return saida

    def _um_contexto(self, actor) -> dict:
        conn = self._connect()
        try:
            cursor = conn.cursor()
            cursor.execute("SELECT chave, nome, ativo FROM dbo.operacoes")
            ops = {}
            for chave, nome, ativo in cursor.fetchall():
                registro = {"chave": normalize_text(chave), "nome": normalize_text(nome), "ativo": bool(ativo)}
                ops[_norm(chave)] = registro
                ops.setdefault(_norm(nome), registro)
            cursor.execute("SELECT LOWER(email), LOWER(login) FROM dbo.usuarios")
            existentes = set()
            for email, login in cursor.fetchall():
                existentes.update(x for x in (normalize_text(email), normalize_text(login)) if x)
            cursor.execute("SELECT id_usuario, email, nome, sobrenome, perfil_id, status FROM dbo.usuarios")
            supervisores = {}
            for id_usuario, email, nome, sobrenome, perfil, st in cursor.fetchall():
                if normalize_text(perfil) == ROLE_SUPERVISOR and normalize_text(st) == "Ativo" and normalize_text(email):
                    supervisores[normalize_text(email).lower()] = {"id": int(id_usuario), "nome": f"{normalize_text(nome)} {normalize_text(sobrenome)}".strip(), "operacoes": set()}
            cursor.execute("SELECT id_usuario, operacao FROM dbo.usuarios_operacoes")
            por_id = {s["id"]: s for s in supervisores.values()}
            for id_usuario, operacao in cursor.fetchall():
                if int(id_usuario) in por_id:
                    por_id[int(id_usuario)]["operacoes"].add(normalize_text(operacao))
            cursor.execute("SELECT id_equipe, operacao, nome, ativo FROM dbo.equipes_operacao")
            equipes = [{"id": int(r[0]), "operacao": normalize_text(r[1]), "nome": normalize_text(r[2]), "ativo": bool(r[3])} for r in cursor.fetchall()]
            cursor.execute("SELECT id_item, operacao, valor, ativo FROM dbo.monitoria_catalogo WHERE tipo = 'canal'")
            canais = [{"id": int(r[0]), "operacao": normalize_text(r[1]), "nome": normalize_text(r[2]), "ativo": bool(r[3])} for r in cursor.fetchall()]
        finally:
            conn.close()
        perfis = {}
        for r in _perfis_visiveis():
            perfis[_norm(r.name)] = r
            perfis.setdefault(_norm(r.id), r)
        return {
            "operacoes": ops, "existentes": existentes, "supervisores": supervisores, "equipes": equipes, "canais": canais, "perfis": perfis,
            "turnos": {_norm(t["valor"]): t["valor"] for t in self.mon_list_catalogo("turno")},
        }

    def _um_analisar_linha(self, actor, ctx: dict, dados: dict, vistos: set) -> tuple[dict | None, str]:
        """(usuário pronto para criar, '') ou (None, motivo da rejeição)."""
        for rotulo, chave, obrigatoria, _ in COLUNAS:
            if obrigatoria and not dados[chave]:
                return None, f"Campo obrigatório vazio: {rotulo}."
        email = dados["email"].lower()
        if not _EMAIL.match(email) or len(email) > 180:
            return None, "E-mail inválido."
        if email in ctx["existentes"]:
            return None, "Já existe um usuário com este e-mail."
        if email in vistos:
            return None, "E-mail repetido na planilha (já usado em outra linha)."
        papel = ctx["perfis"].get(_norm(dados["perfil"]))
        if not papel:
            return None, f"Perfil não reconhecido: {dados['perfil']}."
        if papel.id == ROLE_ADMIN and actor.perfil != ROLE_ADMIN:
            return None, "Somente o Administrador cria usuários Administradores."
        if papel.id in _PERFIS_MONITORIA and actor.perfil != ROLE_ADMIN and not pode_gerenciar_perfil(actor.perfil, papel.id):
            return None, "Seu perfil não pode criar este nível de usuário."
        acesso = {"microsoft": "microsoft", "local": "local"}.get(_norm(dados["acesso"]))
        if not acesso:
            return None, "Tipo de acesso deve ser Microsoft ou Local."
        senha = dados["senha"]
        if acesso == "local":
            if len(senha) < 8:
                return None, "Acesso Local exige senha inicial com pelo menos 8 caracteres."
        else:
            senha = ""

        operacoes = []
        for item in _lista(dados["operacao"]):
            op = ctx["operacoes"].get(_norm(item))
            if not op:
                return None, f"Operação não encontrada: {item}."
            if not op["ativo"]:
                return None, f"Operação inativa: {op['nome']}."
            if not pode_ver_operacao(actor.perfil, actor.operacoes, op["chave"]):
                return None, f"Operação fora do seu escopo: {op['nome']}."
            if op["chave"] not in operacoes:
                operacoes.append(op["chave"])

        supervisores = []
        for item in _lista(dados["supervisor"]):
            sup = ctx["supervisores"].get(item.lower())
            if not sup:
                return None, f"Supervisor não encontrado (use o e-mail de um supervisor ativo): {item}."
            if sup["id"] not in [s["id"] for s in supervisores]:
                supervisores.append(sup)
        id_equipe, turno, canais = None, None, []
        if papel.id in _PERFIS_MONITORIA:
            erros = validar_vinculos(papel.id, operacoes, [s["id"] for s in supervisores])
            if erros:
                return None, " ".join(erros)
            for sup in supervisores:
                if not (sup["operacoes"] & set(operacoes)):
                    return None, f"O supervisor {sup['nome']} não está vinculado à operação informada."
            if dados["equipe"]:
                achada = [e for e in ctx["equipes"] if e["ativo"] and e["operacao"] in operacoes and _norm(e["nome"]) == _norm(dados["equipe"])]
                if papel.id != ROLE_OPERATOR or not achada:
                    return None, f"Equipe não encontrada na operação (só para Operador): {dados['equipe']}."
                id_equipe = achada[0]["id"]
            if dados["turno"]:
                turno = ctx["turnos"].get(_norm(dados["turno"]))
                if not turno or papel.id not in (ROLE_OPERATOR, ROLE_SUPERVISOR):
                    return None, f"Turno inválido (Operador/Supervisor): {dados['turno']}."
            for item in _lista(dados["canais"]):
                achado = [c for c in ctx["canais"] if c["ativo"] and c["operacao"] in operacoes and _norm(c["nome"]) == _norm(item)]
                if not achado or papel.id not in (ROLE_OPERATOR, ROLE_SUPERVISOR):
                    return None, f"Canal não encontrado na operação (Operador/Supervisor): {item}."
                if achado[0]["id"] not in canais:
                    canais.append(achado[0]["id"])
        else:
            if supervisores:
                return None, "Somente o Operador possui supervisor responsável."
            if dados["equipe"] or dados["turno"] or dados["canais"]:
                return None, "Equipe, turno e canais só se aplicam a Operador e Supervisor."
        vistos.add(email)
        return {
            "nome": dados["nome"], "sobrenome": dados["sobrenome"], "email": email, "login": email, "cargo": dados["cargo"],
            "perfil": papel.id, "perfil_nome": papel.name, "provedor_autenticacao": acesso, "senha": senha,
            "operacoes": operacoes, "supervisores": [s["id"] for s in supervisores], "supervisores_nomes": [s["nome"] for s in supervisores],
            "id_equipe": id_equipe, "turno": turno, "canais": canais,
        }, ""

    def usuarios_massa_processar(self, actor, conteudo: bytes, *, confirmar: bool = False, ip: str = "") -> dict:
        linhas = self._um_ler_planilha(conteudo)
        ctx = self._um_contexto(actor)
        vistos: set[str] = set()
        validos, ignorados = [], []
        for numero, dados in linhas:
            usuario, motivo = self._um_analisar_linha(actor, ctx, dados, vistos)
            rotulo = " ".join(p for p in (dados["nome"], dados["sobrenome"]) if p) or "(sem nome)"
            if usuario is None:
                ignorados.append({"linha": numero, "nome": rotulo, "email": dados["email"], "motivo": motivo})
            else:
                validos.append((numero, usuario))

        def resumo(numero: int, u: dict) -> dict:
            nomes = {o["chave"]: o["nome"] for o in ctx["operacoes"].values()}
            return {
                "linha": numero, "nome": f"{u['nome']} {u['sobrenome']}".strip(), "email": u["email"], "cargo": u["cargo"],
                "perfil": u["perfil_nome"], "acesso": "Local" if u["provedor_autenticacao"] == "local" else "Microsoft",
                "operacoes": [nomes.get(c, c) for c in u["operacoes"]], "supervisores": u["supervisores_nomes"],
            }

        resultado = {
            "confirmado": bool(confirmar),
            "total_linhas": len(linhas),
            "validos": [resumo(n, u) for n, u in validos],
            "ignorados": ignorados,
        }
        if not confirmar:
            return resultado

        criados, falhas = [], []
        for numero, u in validos:
            try:
                if u["perfil"] in _PERFIS_MONITORIA:
                    self.mon_create_usuario(actor, {**u, "supervisores": u["supervisores"]}, ip=ip)
                else:
                    criado = self.create_system_user(
                        {k: u[k] for k in ("nome", "sobrenome", "email", "login", "cargo", "perfil", "provedor_autenticacao", "senha", "operacoes")} | {"status": "Ativo"},
                        actor=actor,
                    )
                    if u["provedor_autenticacao"] == "local":
                        conn = self._connect()
                        try:
                            conn.cursor().execute("UPDATE dbo.usuarios SET deve_trocar_senha = 1 WHERE id_usuario = ?", (int(criado["id_usuario"]),))
                            conn.commit()
                        finally:
                            conn.close()
                criados.append(resumo(numero, u))
            except HTTPException as exc:
                falhas.append({"linha": numero, "nome": f"{u['nome']} {u['sobrenome']}".strip(), "email": u["email"], "motivo": str(exc.detail)})
            except Exception:  # noqa: BLE001 - uma linha com erro inesperado não derruba as demais
                falhas.append({"linha": numero, "nome": f"{u['nome']} {u['sobrenome']}".strip(), "email": u["email"], "motivo": "Erro inesperado ao criar este usuário."})
        resultado["criados"] = criados
        resultado["falhas"] = falhas
        return resultado
